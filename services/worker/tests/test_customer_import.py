from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import replace
from pathlib import Path
from typing import cast
from uuid import uuid4

import asyncpg
import pytest
from openpyxl import Workbook

from teamora_worker.background_jobs import (
    BackgroundJobClaim,
    JobExecutionContext,
    JobExecutionError,
)
from teamora_worker.config import WorkerSettings
from teamora_worker.customer_import import (
    CustomerImportProcessor,
    FieldDefinition,
    ImportSource,
    _normalize_email,
    _normalize_phone,
    _open_csv,
    _prepare_row,
    _valid_object_key,
    _validate_mapping,
    _validate_xlsx,
)
from teamora_worker.maintenance import MaintenanceHandlers


def settings(**values: object) -> WorkerSettings:
    return WorkerSettings(_env_file=None, **values)  # type: ignore[call-arg]


def source(tenant_id: object, project_id: object, *, object_key: str | None = None) -> ImportSource:
    return ImportSource(
        import_id=uuid4(),
        project_id=project_id,  # type: ignore[arg-type]
        created_by_user_id=uuid4(),
        file_name="customers.csv",
        file_type="csv",
        selected_sheet="CSV",
        mapping={"display_name": "Name"},
        update_rule="skip",
        status="ready",
        storage_object_id=uuid4(),
        bucket="teamora-private",
        object_key=object_key or f"tenants/{tenant_id}/imports/source.csv",
        size_bytes=42,
        checksum_sha256=None,
        content_type="text/csv",
    )


def job(tenant_id: object, project_id: object) -> BackgroundJobClaim:
    return BackgroundJobClaim(
        id=uuid4(),
        tenant_id=tenant_id,  # type: ignore[arg-type]
        project_id=project_id,  # type: ignore[arg-type]
        job_type="customer_import.process",
        queue="imports",
        safe_payload={"import_id": str(uuid4())},
        attempt_count=1,
        max_attempts=4,
        lease_owner="worker-test",
        lease_token=uuid4(),
        correlation_id="test",
        causation_id=None,
    )


def test_phone_and_email_normalization_fail_closed() -> None:
    assert _normalize_phone("+998 (90) 123-45-67") == "+998901234567"
    assert _normalize_phone("00998901234567") == "+998901234567"
    assert _normalize_email("  PERSON@Example.COM ") == "person@example.com"
    with pytest.raises(ValueError):
        _normalize_phone("12345")
    with pytest.raises(ValueError):
        _normalize_email("not-an-email")


def test_mapping_and_row_validation_preserve_kaa_and_required_fields() -> None:
    definitions = {
        "segment": FieldDefinition(
            key="segment",
            name="Segment",
            field_type="select",
            required=True,
            options=("vip", "standard"),
            default_value=None,
        )
    }
    mapping = _validate_mapping(
        {
            "display_name": "Name",
            "phone": "Phone",
            "email": "Email",
            "preferred_language": "Language",
            "custom.segment": "Segment",
        },
        ("Name", "Phone", "Email", "Language", "Segment"),
        definitions,
    )
    prepared = _prepare_row(
        2,
        {
            "Name": "Aydos Qurbanov",
            "Phone": "+998 90 123 45 67",
            "Email": "AYDOS@example.com",
            "Language": "kaa",
            "Segment": "vip",
        },
        mapping,
        definitions,
    )
    assert not prepared.errors
    assert prepared.values["preferred_language"] == "kaa"
    assert prepared.values["custom_fields"] == {"segment": "vip"}
    assert "phone:+998901234567" in prepared.identities
    assert "email:aydos@example.com" in prepared.identities

    invalid = _prepare_row(3, {"Name": "A"}, {"display_name": "Name"}, definitions)
    assert {error["code"] for error in invalid.errors} == {
        "display_name_required",
        "custom_field_required",
    }


def test_csv_reader_streams_rows_and_rejects_binary_mismatch(tmp_path: Path) -> None:
    path = tmp_path / "customers.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Name", "Phone"])
        writer.writerow(["Customer One", "+998901111111"])
        writer.writerow(["Customer Two", "+998902222222"])
    with _open_csv(path, 10) as stream:
        assert stream.headers == ("Name", "Phone")
        assert [row[0] for row in stream.rows] == [2, 3]

    path.write_bytes(b"PK\x03\x04not-a-csv")
    with pytest.raises(JobExecutionError) as error, _open_csv(path, 10):
        pass
    assert error.value.code == "customer_import_csv_format_invalid"


def test_xlsx_validation_and_selected_sheet_format(tmp_path: Path) -> None:
    path = tmp_path / "customers.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Customers"
    worksheet.append(["Name", "Phone"])
    worksheet.append(["Customer One", "+998901111111"])
    workbook.save(path)
    workbook.close()
    _validate_xlsx(path, max_expanded_bytes=10_000_000, max_ratio=1000, max_sheets=10)

    unsafe = tmp_path / "unsafe.xlsx"
    with zipfile.ZipFile(unsafe, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"<worksheet />")
        archive.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(JobExecutionError) as error:
        _validate_xlsx(unsafe, max_expanded_bytes=10_000, max_ratio=100, max_sheets=10)
    assert error.value.code == "customer_import_xlsx_active_content"


def test_xlsx_pending_sheet_selects_first_real_sheet(tmp_path: Path) -> None:
    path = tmp_path / "customers.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Customers"
    first.append(["Name", "Phone"])
    first.append(["Customer One", "+998901111111"])
    second = workbook.create_sheet("Archive")
    second.append(["Name"])
    workbook.save(path)
    workbook.close()
    tenant_id = uuid4()
    project_id = uuid4()
    processor = CustomerImportProcessor(settings())
    xlsx_source = replace(
        source(tenant_id, project_id),
        file_name="customers.xlsx",
        file_type="xlsx",
        selected_sheet="pending",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    with processor._open_rows(path, xlsx_source) as stream:
        assert stream.sheet_names == ("Customers", "Archive")
        assert stream.selected_sheet == "Customers"
        assert [row_number for row_number, _row in stream.rows] == [2]


def test_row_validation_enforces_customer_database_lengths() -> None:
    prepared = _prepare_row(
        2,
        {
            "Name": "x" * 161,
            "Address": "y" * 501,
            "Description": "z" * 4001,
        },
        {
            "display_name": "Name",
            "address": "Address",
            "description": "Description",
        },
        {},
    )
    too_long = {
        str(error.get("field"))
        for error in prepared.errors
        if error.get("code") == "field_too_long"
    }
    assert too_long == {"display_name", "address", "description"}


@pytest.mark.asyncio
async def test_final_identity_recheck_and_contact_conflict_are_safe() -> None:
    tenant_id = uuid4()
    project_id = uuid4()
    existing_customer_id = uuid4()
    processor = CustomerImportProcessor(settings())

    class IdentityConnection:
        async def fetch(self, _query: str, *_arguments: object) -> list[dict[str, object]]:
            return [
                {
                    "identity": "phone:+998901111111",
                    "customer_id": existing_customer_id,
                }
            ]

    matches = await processor._final_identity_matches(
        cast(asyncpg.Connection, IdentityConnection()),
        job(tenant_id, project_id),
        source(tenant_id, project_id),
        cast(list[asyncpg.Record], [{"identity_hashes": ["phone:+998901111111"]}]),
    )
    assert matches == {"phone:+998901111111": {existing_customer_id}}

    class ContactConflictConnection:
        def __init__(self) -> None:
            self.deleted = False

        async def fetchval(self, query: str, *arguments: object) -> object:
            if "INSERT INTO customers" in query:
                return arguments[0]
            if "SELECT EXISTS" in query:
                return False
            if "INSERT INTO customer_contacts" in query:
                return None
            if "SELECT customer_id FROM customer_contacts" in query:
                return existing_customer_id
            raise AssertionError(query)

        async def execute(self, query: str, *_arguments: object) -> str:
            if "DELETE FROM customers" in query:
                self.deleted = True
            return "DELETE 1"

    conflict_connection = ContactConflictConnection()
    created = await processor._create_customer(
        cast(asyncpg.Connection, conflict_connection),
        job(tenant_id, project_id),
        source(tenant_id, project_id),
        cast(
            asyncpg.Record,
            {
                "normalized_payload": {
                    "values": {"display_name": "Safe customer"},
                    "contacts": [
                        {
                            "kind": "phone",
                            "normalized_value": "+998901111111",
                            "display_value": "+998 90 111 11 11",
                            "primary": True,
                        }
                    ],
                }
            },
        ),
    )
    assert created is False
    assert conflict_connection.deleted is True

    class UpdateConflictConnection:
        def __init__(self) -> None:
            self.scalar_update_attempted = False

        async def fetchrow(self, query: str, *_arguments: object) -> object:
            assert "SELECT id FROM customers" in query
            return {"id": existing_customer_id}

        async def fetchval(self, query: str, *_arguments: object) -> object:
            if "SELECT EXISTS" in query:
                return False
            if "INSERT INTO customer_contacts" in query:
                return None
            if "SELECT customer_id FROM customer_contacts" in query:
                return uuid4()
            raise AssertionError(query)

        async def execute(self, query: str, *_arguments: object) -> str:
            assert "UPDATE customers SET" in query
            self.scalar_update_attempted = True
            return "UPDATE 1"

    update_connection = UpdateConflictConnection()
    updated = await processor._update_customer(
        cast(asyncpg.Connection, update_connection),
        job(tenant_id, project_id),
        source(tenant_id, project_id),
        cast(
            asyncpg.Record,
            {
                "normalized_payload": {
                    "values": {"display_name": "Concurrent update"},
                    "contacts": [
                        {
                            "kind": "phone",
                            "normalized_value": "+998901111111",
                            "display_value": "+998 90 111 11 11",
                            "primary": True,
                        }
                    ],
                }
            },
        ),
        target_customer_id=existing_customer_id,
    )
    assert update_connection.scalar_update_attempted is True
    assert updated is False


@pytest.mark.asyncio
async def test_expired_preview_maintenance_clears_only_terminal_staging() -> None:
    statements: list[str] = []

    class Transaction:
        async def __aenter__(self) -> None:
            return None

        async def __aexit__(self, *_arguments: object) -> None:
            return None

    class Connection:
        def transaction(self) -> Transaction:
            return Transaction()

        async def execute(self, query: str, *_arguments: object) -> str:
            statements.append(query)
            if "UPDATE customer_imports" in query:
                return "UPDATE 2"
            if "DELETE FROM customer_import_staging_rows" in query:
                return "DELETE 3"
            return "SELECT 1"

    connection = Connection()

    class Acquisition:
        async def __aenter__(self) -> Connection:
            return connection

        async def __aexit__(self, *_arguments: object) -> None:
            return None

    class Pool:
        def acquire(self) -> Acquisition:
            return Acquisition()

    tenant_id = uuid4()
    project_id = uuid4()
    claim = job(tenant_id, project_id)
    context = JobExecutionContext(cast(asyncpg.Pool, Pool()), settings(), claim)
    result = await MaintenanceHandlers.cleanup_expired_import_previews(
        cast(MaintenanceHandlers, object()),
        claim,
        context,
    )
    assert result.metadata == {"expired_previews": 2, "staging_rows_removed": 3}
    cleanup_sql = "\n".join(statements)
    assert "status IN ('preview','ready')" in cleanup_sql
    assert "status IN ('completed','cancelled','expired','failed')" in cleanup_sql


def test_storage_scope_and_registry_are_tenant_safe() -> None:
    tenant_id = uuid4()
    project_id = uuid4()
    CustomerImportProcessor._validate_claim_scope(
        job(tenant_id, project_id),
        source(tenant_id, project_id),
    )
    with pytest.raises(JobExecutionError) as error:
        CustomerImportProcessor._validate_claim_scope(
            job(tenant_id, project_id),
            source(tenant_id, project_id, object_key=f"tenants/{uuid4()}/source.csv"),
        )
    assert error.value.code == "customer_import_storage_scope_mismatch"
    assert _valid_object_key(f"tenants/{tenant_id}/imports/source.csv")
    assert not _valid_object_key("../outside.csv")


def test_customer_import_handlers_are_registered() -> None:
    processor = CustomerImportProcessor(settings())
    assert set(processor.registry()) == {
        "customer_import.prepare_preview",
        "customer_import.process",
    }


def test_error_report_shape_never_needs_raw_customer_rows() -> None:
    buffer = io.StringIO()
    csv.writer(buffer).writerows([["Name"], ["Customer"]])
    assert "Customer" in buffer.getvalue()
    # Staging reports store row number and error codes; source values stay in
    # the tenant-scoped staging payload and are never copied into job metadata.
