from __future__ import annotations

import io
import math
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

import teamora_worker.knowledge_ingestion as ingestion
from teamora_worker.knowledge_ingestion import (
    DocumentProcessingError,
    NeedsOcrError,
    chunk_blocks,
    deterministic_embedding,
    extract_document,
)


def test_txt_markdown_and_docx_extraction_preserve_unicode_and_sections() -> None:
    text = "# Qaraqalpaq\n\nQollap-quwatlaw xizmeti hár kúni isleydi."
    extracted = extract_document(text.encode(), "md", max_pdf_pages=500, max_chars=10_000)
    assert extracted.metadata["format"] == "md"
    assert extracted.blocks[0].section == "Qaraqalpaq"
    assert "hár kúni" in extracted.normalized_text

    document = Document()
    document.add_heading("Support", level=1)
    document.add_paragraph("The service works every weekday.")
    stream = io.BytesIO()
    document.save(stream)
    docx = extract_document(stream.getvalue(), "docx", max_pdf_pages=500, max_chars=10_000)
    assert docx.metadata["format"] == "docx"
    assert docx.blocks[-1].section == "Support"


def test_pdf_text_layer_page_metadata_and_scanned_detection(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class Page:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class Reader:
        is_encrypted = False
        pages = [Page("Support is available Monday through Friday."), Page("Second page source.")]

        def __init__(self, _stream: object, strict: bool) -> None:
            assert strict is True

    monkeypatch.setattr(ingestion, "PdfReader", Reader)
    extracted = extract_document(b"%PDF-test", "pdf", max_pdf_pages=500, max_chars=10_000)
    assert extracted.page_count == 2
    assert [block.page for block in extracted.blocks] == [1, 2]

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = io.BytesIO()
    writer.write(stream)
    monkeypatch.undo()
    with pytest.raises(NeedsOcrError):
        extract_document(stream.getvalue(), "pdf", max_pdf_pages=500, max_chars=10_000)


def test_encrypted_corrupted_and_page_limited_pdf() -> None:
    with pytest.raises(DocumentProcessingError) as corrupt:
        extract_document(b"not-pdf", "pdf", max_pdf_pages=500, max_chars=10_000)
    assert corrupt.value.code == "pdf_corrupt"

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("test-password")
    encrypted = io.BytesIO()
    writer.write(encrypted)
    with pytest.raises(DocumentProcessingError) as error:
        extract_document(encrypted.getvalue(), "pdf", max_pdf_pages=500, max_chars=10_000)
    assert error.value.code == "pdf_encrypted"

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_blank_page(width=100, height=100)
    pages = io.BytesIO()
    writer.write(pages)
    with pytest.raises(DocumentProcessingError) as limit:
        extract_document(pages.getvalue(), "pdf", max_pdf_pages=1, max_chars=10_000)
    assert limit.value.code == "pdf_page_limit"


def test_docx_macro_traversal_external_relationship_and_zip_bomb() -> None:
    def archive(entries: dict[str, bytes]) -> bytes:
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as value:
            for name, content in entries.items():
                value.writestr(name, content)
        return stream.getvalue()

    base = {"[Content_Types].xml": b"types", "word/document.xml": b"document"}
    cases = [
        ({**base, "word/vbaProject.bin": b"macro"}, "docx_macro"),
        ({**base, "../outside": b"unsafe"}, "path_traversal"),
        (
            {
                **base,
                "word/_rels/document.xml.rels": b'<Relationship TargetMode="External" />',
            },
            "docx_external_relationship",
        ),
    ]
    for entries, code in cases:
        with pytest.raises(DocumentProcessingError) as error:
            extract_document(archive(entries), "docx", max_pdf_pages=500, max_chars=10_000)
        assert error.value.code == code

    with pytest.raises(DocumentProcessingError) as bomb:
        extract_document(
            archive({**base, "word/large.xml": b"x" * 20_000}),
            "docx",
            max_pdf_pages=500,
            max_chars=10_000,
            max_docx_expanded_bytes=1_000,
        )
    assert bomb.value.code == "zip_bomb"


def test_chunking_and_mock_embeddings_are_deterministic_but_not_production_claims() -> None:
    blocks = [
        ingestion.ExtractedBlock(
            "First sentence contains the exact support schedule. "
            "Second sentence keeps the source readable. " * 10,
            page=3,
            section="Schedule",
        )
    ]
    first = chunk_blocks(blocks, size_chars=240, overlap_chars=30)
    second = chunk_blocks(blocks, size_chars=240, overlap_chars=30)
    assert [chunk.checksum for chunk in first] == [chunk.checksum for chunk in second]
    assert all(chunk.page == 3 and chunk.section == "Schedule" for chunk in first)
    assert all(chunk.character_count <= 240 for chunk in first)

    vector = deterministic_embedding("support schedule weekday", 64)
    repeated = deterministic_embedding("support schedule weekday", 64)
    unrelated = deterministic_embedding("orchid satellite", 64)
    assert vector == repeated
    assert math.isclose(sum(value * value for value in vector), 1.0)
    assert sum(left * right for left, right in zip(vector, unrelated, strict=True)) < 0.8


def test_limits_and_invalid_chunk_config_fail_closed() -> None:
    with pytest.raises(DocumentProcessingError) as text_limit:
        extract_document(b"x" * 101, "txt", max_pdf_pages=500, max_chars=100)
    assert text_limit.value.code == "extracted_text_limit"
    with pytest.raises(DocumentProcessingError) as chunk_config:
        chunk_blocks([ingestion.ExtractedBlock("usable text")], size_chars=200, overlap_chars=200)
    assert chunk_config.value.code == "chunk_config_invalid"


def test_asyncpg_json_configuration_decodes_deterministically() -> None:
    assert ingestion.json_object('{"size_chars": 1200, "overlap_chars": 120}') == {
        "size_chars": 1200,
        "overlap_chars": 120,
    }
    assert ingestion.json_object({"size_chars": 800}) == {"size_chars": 800}
    assert ingestion.json_object([]) == {}
