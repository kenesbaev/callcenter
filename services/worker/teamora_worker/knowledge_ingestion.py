from __future__ import annotations

import asyncio
import hashlib
import io
import json
import math
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath
from uuid import UUID, uuid4

import asyncpg
import structlog
from docx import Document
from minio import Minio
from pypdf import PdfReader

from teamora_worker.config import WorkerSettings

log = structlog.get_logger(service="worker", component="knowledge_ingestion")


class DocumentProcessingError(RuntimeError):
    def __init__(self, code: str, message: str, *, transient: bool = False) -> None:
        super().__init__(message)
        self.code = code
        self.safe_message = message[:500]
        self.transient = transient


class NeedsOcrError(DocumentProcessingError):
    def __init__(self) -> None:
        super().__init__("needs_ocr", "PDF has no usable text layer")


@dataclass(frozen=True)
class ExtractedBlock:
    text: str
    page: int | None = None
    section: str | None = None


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    normalized_text: str
    blocks: list[ExtractedBlock]
    page_count: int
    metadata: dict[str, object]


@dataclass(frozen=True)
class ChunkDraft:
    text: str
    page: int | None
    section: str | None
    checksum: str
    character_count: int
    token_count: int


@dataclass(frozen=True)
class ClaimedJob:
    id: UUID
    tenant_id: UUID
    project_id: UUID
    document_version_id: UUID
    lease_token: UUID
    attempts: int
    max_attempts: int
    object_key: str
    file_type: str
    checksum_sha256: str
    language_code: str
    document_id: UUID
    revision_id: UUID
    embedding_provider: str
    embedding_model: str
    embedding_dimension: int
    index_version: str
    chunking_config: dict[str, object]


def normalize_text(value: str) -> str:
    value = value.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.split("\n")]
    return "\n".join(lines).strip()


def extract_document(
    data: bytes,
    file_type: str,
    *,
    max_pdf_pages: int,
    max_chars: int,
    max_docx_expanded_bytes: int = 200 * 1024 * 1024,
    max_zip_ratio: int = 100,
) -> ExtractedDocument:
    if file_type in {"txt", "md"}:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise DocumentProcessingError("text_encoding", "Text documents must use UTF-8") from exc
        blocks = _text_blocks(text)
        page_count = 1
        metadata: dict[str, object] = {"format": file_type}
    elif file_type == "pdf":
        try:
            reader = PdfReader(io.BytesIO(data), strict=True)
        except Exception as exc:
            raise DocumentProcessingError("pdf_corrupt", "PDF is corrupted or unsupported") from exc
        if reader.is_encrypted:
            raise DocumentProcessingError(
                "pdf_encrypted", "Encrypted PDF documents are not supported"
            )
        page_count = len(reader.pages)
        if page_count > max_pdf_pages:
            raise DocumentProcessingError("pdf_page_limit", "PDF exceeds the configured page limit")
        blocks = []
        for index, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                raise DocumentProcessingError(
                    "pdf_extract_failed", "PDF text extraction failed"
                ) from exc
            if text.strip():
                blocks.append(ExtractedBlock(text=text, page=index))
        text = "\n\n".join(block.text for block in blocks)
        if len(normalize_text(text)) < max(20, page_count * 10):
            raise NeedsOcrError()
        metadata = {"format": "pdf", "pages": page_count}
    elif file_type == "docx":
        _validate_docx(
            data,
            max_expanded_bytes=max_docx_expanded_bytes,
            max_zip_ratio=max_zip_ratio,
        )
        try:
            document = Document(io.BytesIO(data))
        except Exception as exc:
            raise DocumentProcessingError("docx_corrupt", "DOCX could not be parsed") from exc
        blocks = []
        section: str | None = None
        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if not value:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                section = value[:500]
            blocks.append(ExtractedBlock(text=value, section=section))
        text = "\n\n".join(block.text for block in blocks)
        page_count = 1
        metadata = {"format": "docx", "paragraphs": len(blocks)}
    else:
        raise DocumentProcessingError("file_type_unsupported", "Unsupported knowledge file type")
    normalized = normalize_text(text)
    if not normalized:
        raise DocumentProcessingError("document_empty", "Document contains no usable text")
    if len(normalized) > max_chars:
        raise DocumentProcessingError(
            "extracted_text_limit", "Extracted text exceeds the configured limit"
        )
    return ExtractedDocument(
        text=text,
        normalized_text=normalized,
        blocks=blocks,
        page_count=page_count,
        metadata=metadata,
    )


def _text_blocks(text: str) -> list[ExtractedBlock]:
    result: list[ExtractedBlock] = []
    section: str | None = None
    for paragraph in (part.strip() for part in re.split(r"\n\s*\n", text)):
        if not paragraph:
            continue
        if paragraph.startswith("#"):
            section = paragraph.lstrip("# ")[:500]
        result.append(ExtractedBlock(paragraph, page=1, section=section))
    return result


def _validate_docx(data: bytes, *, max_expanded_bytes: int, max_zip_ratio: int) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            expanded = sum(item.file_size for item in entries)
            compressed = max(1, sum(item.compress_size for item in entries))
            if (
                len(entries) > 10_000
                or expanded > max_expanded_bytes
                or expanded / compressed > max_zip_ratio
            ):
                raise DocumentProcessingError("zip_bomb", "DOCX expansion is unsafe")
            names = {item.filename.replace("\\", "/") for item in entries}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocumentProcessingError("docx_structure", "DOCX structure is incomplete")
            for name in names:
                path = PurePosixPath(name)
                if path.is_absolute() or ".." in path.parts:
                    raise DocumentProcessingError("path_traversal", "DOCX contains an unsafe path")
                if name.lower().endswith(("vbaproject.bin", ".exe", ".dll", ".js")):
                    raise DocumentProcessingError("docx_macro", "Macro-enabled DOCX is forbidden")
            for relation in (name for name in names if name.lower().endswith(".rels")):
                content = archive.read(relation)
                if b'TargetMode="External"' in content or b"TargetMode='External'" in content:
                    raise DocumentProcessingError(
                        "docx_external_relationship", "External DOCX relationships are forbidden"
                    )
    except zipfile.BadZipFile as exc:
        raise DocumentProcessingError("docx_corrupt", "DOCX archive is corrupted") from exc


def chunk_blocks(
    blocks: list[ExtractedBlock], *, size_chars: int, overlap_chars: int
) -> list[ChunkDraft]:
    if size_chars < 200 or overlap_chars < 0 or overlap_chars >= size_chars:
        raise DocumentProcessingError("chunk_config_invalid", "Chunking configuration is invalid")
    chunks: list[ChunkDraft] = []
    for block in blocks:
        value = normalize_text(block.text)
        start = 0
        while start < len(value):
            end = min(start + size_chars, len(value))
            if end < len(value):
                boundary = max(
                    value.rfind(mark, start + size_chars // 2, end)
                    for mark in (". ", "! ", "? ", "\n", " ")
                )
                if boundary > start:
                    end = boundary + 1
            piece = value[start:end].strip()
            if piece:
                chunks.append(
                    ChunkDraft(
                        text=piece,
                        page=block.page,
                        section=block.section,
                        checksum=hashlib.sha256(piece.encode()).hexdigest(),
                        character_count=len(piece),
                        token_count=max(1, len(re.findall(r"\S+", piece))),
                    )
                )
            if end >= len(value):
                break
            start = max(end - overlap_chars, start + 1)
    return chunks


def deterministic_embedding(text: str, dimension: int) -> list[float]:
    vector = [0.0] * dimension
    for token in re.findall(r"[\w'-]+", text.casefold(), flags=re.UNICODE):
        digest = hashlib.blake2b(token.encode(), digest_size=16).digest()
        vector[int.from_bytes(digest[:8], "big") % dimension] += 1.0 if digest[8] & 1 else -1.0
    norm = math.sqrt(sum(item * item for item in vector))
    return [item / norm for item in vector] if norm else vector


def json_object(value: object) -> dict[str, object]:
    if isinstance(value, str):
        parsed = json.loads(value)
        return dict(parsed) if isinstance(parsed, dict) else {}
    return dict(value) if isinstance(value, dict) else {}


class KnowledgeIngestionProcessor:
    def __init__(self, settings: WorkerSettings) -> None:
        self.settings = settings
        endpoint = settings.minio_endpoint.removeprefix("http://").removeprefix("https://")
        self.storage = Minio(
            endpoint,
            access_key=settings.minio_root_user,
            secret_key=settings.minio_root_password,
            secure=settings.minio_endpoint.startswith("https://"),
        )

    async def run(self, stopping: asyncio.Event) -> None:
        connection: asyncpg.Connection | None = None
        while not stopping.is_set():
            try:
                if connection is None or connection.is_closed():
                    connection = await asyncpg.connect(
                        self.settings.asyncpg_database_url, timeout=5
                    )
                job = await self.claim_once(connection)
                if job is None:
                    try:
                        await asyncio.wait_for(
                            stopping.wait(), self.settings.knowledge_job_poll_seconds
                        )
                    except TimeoutError:
                        pass
                    continue
                await self.process(connection, job)
            except asyncio.CancelledError:
                raise
            except Exception as exc:
                log.warning(
                    "knowledge_worker_retry",
                    error_type=type(exc).__name__,
                    database_code=getattr(exc, "sqlstate", None),
                    database_message=getattr(exc, "message", None),
                )
                if connection is not None and not connection.is_closed():
                    await connection.close()
                connection = None
                try:
                    await asyncio.wait_for(stopping.wait(), 1.0)
                except TimeoutError:
                    pass
        if connection is not None and not connection.is_closed():
            await connection.close()

    async def claim_once(self, connection: asyncpg.Connection) -> ClaimedJob | None:
        tenants = await connection.fetch("SELECT id FROM tenants ORDER BY id")
        for tenant in tenants:
            tenant_id = tenant["id"]
            lease_token = uuid4()
            async with connection.transaction():
                await connection.execute(
                    "SELECT set_config('app.tenant_id', $1, true)", str(tenant_id)
                )
                row = await connection.fetchrow(
                    """
                    SELECT job.id, job.tenant_id, job.project_id, job.document_version_id,
                           job.attempts, job.max_attempts, version.object_key, version.file_type,
                           version.checksum_sha256, version.language_code, version.document_id,
                           version.revision_id, revision.embedding_provider,
                           revision.embedding_model, revision.embedding_dimension,
                           revision.index_version, revision.chunking_config
                    FROM document_ingestion_jobs AS job
                    JOIN knowledge_document_versions AS version
                      ON version.tenant_id = job.tenant_id AND version.id = job.document_version_id
                    JOIN knowledge_base_revisions AS revision
                      ON revision.tenant_id = job.tenant_id AND revision.id = version.revision_id
                    WHERE job.tenant_id = $1
                      AND job.next_attempt_at <= now()
                      AND (
                        job.status = 'pending' OR
                        (job.status = 'processing' AND job.lease_expires_at < now())
                      )
                      AND job.attempts < job.max_attempts
                    ORDER BY job.next_attempt_at, job.created_at
                    FOR UPDATE OF job SKIP LOCKED
                    LIMIT 1
                    """,
                    tenant_id,
                )
                if row is None:
                    continue
                if not row["object_key"]:
                    await connection.execute(
                        """
                        UPDATE document_ingestion_jobs
                        SET status='failed', safe_error_code='original_missing'
                        WHERE id=$1
                        """,
                        row["id"],
                    )
                    continue
                await connection.execute(
                    """
                    UPDATE document_ingestion_jobs
                    SET status='processing', stage='extracting', attempts=attempts+1,
                        lease_token=$2, lease_expires_at=now()+($3::integer * interval '1 second'),
                        heartbeat_at=now(), updated_at=now()
                    WHERE tenant_id=$1 AND id=$4
                    """,
                    tenant_id,
                    lease_token,
                    self.settings.knowledge_job_lease_seconds,
                    row["id"],
                )
                await self._set_version_status(
                    connection, tenant_id, row["document_version_id"], "extracting"
                )
                await self._event(
                    connection,
                    tenant_id,
                    row["project_id"],
                    row["document_version_id"],
                    "knowledge.document_processing",
                    "extracting",
                )
                return ClaimedJob(
                    id=row["id"],
                    tenant_id=tenant_id,
                    project_id=row["project_id"],
                    document_version_id=row["document_version_id"],
                    lease_token=lease_token,
                    attempts=row["attempts"] + 1,
                    max_attempts=row["max_attempts"],
                    object_key=row["object_key"],
                    file_type=row["file_type"],
                    checksum_sha256=row["checksum_sha256"],
                    language_code=row["language_code"],
                    document_id=row["document_id"],
                    revision_id=row["revision_id"],
                    embedding_provider=row["embedding_provider"],
                    embedding_model=row["embedding_model"],
                    embedding_dimension=row["embedding_dimension"],
                    index_version=row["index_version"],
                    chunking_config=json_object(row["chunking_config"]),
                )
        return None

    async def process(self, connection: asyncpg.Connection, job: ClaimedJob) -> None:
        try:
            data = await asyncio.wait_for(
                asyncio.to_thread(self._download, job.object_key),
                timeout=self.settings.knowledge_parser_timeout_seconds,
            )
            if len(data) > self.settings.knowledge_max_file_bytes:
                raise DocumentProcessingError(
                    "file_too_large", "Stored file exceeds the configured limit"
                )
            if hashlib.sha256(data).hexdigest() != job.checksum_sha256:
                raise DocumentProcessingError(
                    "checksum_mismatch", "Stored original checksum does not match"
                )
            extracted = await asyncio.wait_for(
                asyncio.to_thread(
                    extract_document,
                    data,
                    job.file_type,
                    max_pdf_pages=self.settings.knowledge_max_pdf_pages,
                    max_chars=self.settings.knowledge_max_extracted_chars,
                    max_docx_expanded_bytes=self.settings.knowledge_max_docx_expanded_bytes,
                    max_zip_ratio=self.settings.knowledge_max_zip_ratio,
                ),
                timeout=self.settings.knowledge_parser_timeout_seconds,
            )
            await self._advance(connection, job, "chunking")
            chunks = chunk_blocks(
                extracted.blocks,
                size_chars=int(
                    str(
                        job.chunking_config.get(
                            "size_chars", self.settings.knowledge_chunk_size_chars
                        )
                    )
                ),
                overlap_chars=int(
                    str(
                        job.chunking_config.get(
                            "overlap_chars", self.settings.knowledge_chunk_overlap_chars
                        )
                    )
                ),
            )
            if not chunks:
                raise DocumentProcessingError("chunks_empty", "Document produced no usable chunks")
            if len(chunks) > self.settings.knowledge_max_chunks:
                raise DocumentProcessingError(
                    "chunk_limit", "Document exceeds the configured chunk limit"
                )
            await self._advance(connection, job, "embedding")
            if job.embedding_provider != "mock" or self.settings.app_env not in {
                "development",
                "test",
            }:
                raise DocumentProcessingError(
                    "embedding_provider_unavailable",
                    "Embedding provider unavailable — live verification required",
                )
            vectors = [
                deterministic_embedding(chunk.text, job.embedding_dimension) for chunk in chunks
            ]
            await self._save_ready(connection, job, extracted, chunks, vectors)
        except NeedsOcrError as exc:
            await self._finish_special(connection, job, "needs_ocr", exc)
        except DocumentProcessingError as exc:
            await self._fail(connection, job, exc)
        except Exception as exc:
            await self._fail(
                connection,
                job,
                DocumentProcessingError(
                    "processing_error", "Document processing failed", transient=True
                ),
            )
            log.warning(
                "knowledge_document_failed", error_type=type(exc).__name__, job_id=str(job.id)
            )

    def _download(self, key: str) -> bytes:
        response = self.storage.get_object(self.settings.minio_bucket, key)
        try:
            return response.read(self.settings.knowledge_max_file_bytes + 1)
        finally:
            response.close()
            response.release_conn()

    async def _advance(self, connection: asyncpg.Connection, job: ClaimedJob, stage: str) -> None:
        async with connection.transaction():
            await connection.execute(
                "SELECT set_config('app.tenant_id', $1, true)", str(job.tenant_id)
            )
            valid = await connection.fetchval(
                """
                SELECT lease_token=$3 FROM document_ingestion_jobs
                WHERE tenant_id=$1 AND id=$2 FOR UPDATE
                """,
                job.tenant_id,
                job.id,
                job.lease_token,
            )
            if not valid:
                raise DocumentProcessingError(
                    "lease_lost", "Document processing lease was lost", transient=True
                )
            await connection.execute(
                """
                UPDATE document_ingestion_jobs SET stage=$3::varchar, heartbeat_at=now(),
                    lease_expires_at=now()+($4::integer * interval '1 second'), updated_at=now()
                WHERE tenant_id=$1 AND id=$2
                """,
                job.tenant_id,
                job.id,
                stage,
                self.settings.knowledge_job_lease_seconds,
            )
            await self._set_version_status(
                connection, job.tenant_id, job.document_version_id, stage
            )
            await self._event(
                connection,
                job.tenant_id,
                job.project_id,
                job.document_version_id,
                "knowledge.document_processing",
                stage,
            )

    async def _save_ready(
        self,
        connection: asyncpg.Connection,
        job: ClaimedJob,
        extracted: ExtractedDocument,
        chunks: list[ChunkDraft],
        vectors: list[list[float]],
    ) -> None:
        async with connection.transaction():
            await connection.execute(
                "SELECT set_config('app.tenant_id', $1, true)", str(job.tenant_id)
            )
            token = await connection.fetchval(
                """
                SELECT lease_token FROM document_ingestion_jobs
                WHERE tenant_id=$1 AND id=$2 FOR UPDATE
                """,
                job.tenant_id,
                job.id,
            )
            if token != job.lease_token:
                raise DocumentProcessingError(
                    "lease_lost", "Document processing lease was lost", transient=True
                )
            await connection.execute(
                "DELETE FROM knowledge_chunks WHERE tenant_id=$1 AND document_version_id=$2",
                job.tenant_id,
                job.document_version_id,
            )
            for ordinal, (chunk, vector) in enumerate(zip(chunks, vectors, strict=True)):
                vector_literal = "[" + ",".join(f"{value:.9f}" for value in vector) + "]"
                await connection.execute(
                    """
                    INSERT INTO knowledge_chunks
                        (id, tenant_id, project_id, revision_id, document_id, document_version_id,
                         ordinal, content, normalized_content, token_count, character_count,
                         content_checksum, language_code, page_number, section, embedding,
                         embedding_status, embedding_provider, embedding_model, embedding_dimension,
                         index_version, created_at, updated_at)
                    VALUES (gen_random_uuid(),$1,$2,$3,$4,$5,$6,$7,$7,$8,$9,$10,$11,$12,$13,
                            $14::vector,'ready',$15,$16,$17,$18,now(),now())
                    """,
                    job.tenant_id,
                    job.project_id,
                    job.revision_id,
                    job.document_id,
                    job.document_version_id,
                    ordinal,
                    chunk.text,
                    chunk.token_count,
                    chunk.character_count,
                    chunk.checksum,
                    job.language_code,
                    chunk.page,
                    chunk.section,
                    vector_literal,
                    job.embedding_provider,
                    job.embedding_model,
                    job.embedding_dimension,
                    job.index_version,
                )
            await connection.execute(
                """
                UPDATE knowledge_document_versions
                SET status='ready', extracted_text=$3, normalized_text=$4, page_count=$5,
                    extraction_metadata=$6::jsonb, ready_at=now(), safe_error_code=NULL,
                    safe_error_message=NULL, lock_version=lock_version+1, updated_at=now()
                WHERE tenant_id=$1 AND id=$2
                """,
                job.tenant_id,
                job.document_version_id,
                extracted.text,
                extracted.normalized_text,
                extracted.page_count,
                json.dumps(
                    {
                        **extracted.metadata,
                        "embedding_usage": {
                            "input_items": len(chunks),
                            "input_characters": sum(item.character_count for item in chunks),
                            "estimated_tokens": sum(item.token_count for item in chunks),
                        },
                    }
                ),
            )
            await connection.execute(
                """
                UPDATE knowledge_documents
                SET content=$3, content_hash=$4, updated_at=now()
                WHERE tenant_id=$1 AND id=$2
                """,
                job.tenant_id,
                job.document_id,
                extracted.normalized_text,
                job.checksum_sha256,
            )
            await connection.execute(
                """
                UPDATE document_ingestion_jobs
                SET status='succeeded', stage='ready', completed_at=now(),
                    lease_token=NULL, lease_expires_at=NULL, heartbeat_at=now(), updated_at=now()
                WHERE tenant_id=$1 AND id=$2
                """,
                job.tenant_id,
                job.id,
            )
            await self._event(
                connection,
                job.tenant_id,
                job.project_id,
                job.document_version_id,
                "knowledge.document_ready",
                "ready",
            )
            await self._event(
                connection,
                job.tenant_id,
                job.project_id,
                job.revision_id,
                "knowledge.index_updated",
                "ready",
                aggregate_type="knowledge_revision",
            )

    async def _finish_special(
        self,
        connection: asyncpg.Connection,
        job: ClaimedJob,
        status: str,
        error: DocumentProcessingError,
    ) -> None:
        async with connection.transaction():
            await connection.execute(
                "SELECT set_config('app.tenant_id', $1, true)", str(job.tenant_id)
            )
            token = await connection.fetchval(
                """
                SELECT lease_token FROM document_ingestion_jobs
                WHERE tenant_id=$1 AND id=$2 FOR UPDATE
                """,
                job.tenant_id,
                job.id,
            )
            if token != job.lease_token:
                return
            await connection.execute(
                """
                UPDATE knowledge_document_versions
                SET status=$3::varchar, safe_error_code=$4, safe_error_message=$5,
                    failed_at=now(), lock_version=lock_version+1, updated_at=now()
                WHERE tenant_id=$1 AND id=$2
                """,
                job.tenant_id,
                job.document_version_id,
                status,
                error.code,
                error.safe_message,
            )
            await connection.execute(
                """
                UPDATE document_ingestion_jobs
                SET status='succeeded', stage=$3::varchar, completed_at=now(),
                    lease_token=NULL, lease_expires_at=NULL,
                    safe_error_code=$4, updated_at=now()
                WHERE tenant_id=$1 AND id=$2
                """,
                job.tenant_id,
                job.id,
                status,
                error.code,
            )
            await self._event(
                connection,
                job.tenant_id,
                job.project_id,
                job.document_version_id,
                "knowledge.document_needs_ocr",
                status,
            )

    async def _fail(
        self, connection: asyncpg.Connection, job: ClaimedJob, error: DocumentProcessingError
    ) -> None:
        retry = error.transient and job.attempts < job.max_attempts
        async with connection.transaction():
            await connection.execute(
                "SELECT set_config('app.tenant_id', $1, true)", str(job.tenant_id)
            )
            token = await connection.fetchval(
                """
                SELECT lease_token FROM document_ingestion_jobs
                WHERE tenant_id=$1 AND id=$2 FOR UPDATE
                """,
                job.tenant_id,
                job.id,
            )
            if token != job.lease_token:
                return
            await connection.execute(
                """
                UPDATE knowledge_document_versions
                SET status='failed', safe_error_code=$3, safe_error_message=$4,
                    failed_at=now(), lock_version=lock_version+1, updated_at=now()
                WHERE tenant_id=$1 AND id=$2
                """,
                job.tenant_id,
                job.document_version_id,
                error.code,
                error.safe_message,
            )
            await connection.execute(
                """
                UPDATE document_ingestion_jobs
                SET status=$3::varchar, stage='failed', safe_error_code=$4,
                    next_attempt_at=now()+($5::integer * interval '1 second'),
                    lease_token=NULL, lease_expires_at=NULL, updated_at=now()
                WHERE tenant_id=$1 AND id=$2
                """,
                job.tenant_id,
                job.id,
                "pending" if retry else "failed",
                error.code,
                min(60, 2**job.attempts),
            )
            await self._event(
                connection,
                job.tenant_id,
                job.project_id,
                job.document_version_id,
                "knowledge.document_failed",
                "failed",
            )

    async def _set_version_status(
        self, connection: asyncpg.Connection, tenant_id: UUID, version_id: UUID, status: str
    ) -> None:
        if status in {"extracting", "chunking", "embedding"}:
            await connection.execute(
                """
                UPDATE knowledge_document_versions
                SET status=$3::varchar,
                    extraction_started_at=CASE WHEN $3::varchar='extracting' THEN now()
                                               ELSE extraction_started_at END,
                    chunking_started_at=CASE WHEN $3::varchar='chunking' THEN now()
                                             ELSE chunking_started_at END,
                    embedding_started_at=CASE WHEN $3::varchar='embedding' THEN now()
                                              ELSE embedding_started_at END,
                    lock_version=lock_version+1, updated_at=now()
                WHERE tenant_id=$1 AND id=$2
                """,
                tenant_id,
                version_id,
                status,
            )

    async def _event(
        self,
        connection: asyncpg.Connection,
        tenant_id: UUID,
        project_id: UUID,
        aggregate_id: UUID,
        event_type: str,
        status: str,
        *,
        aggregate_type: str = "knowledge_document_version",
    ) -> None:
        await connection.execute(
            """
            INSERT INTO realtime_events
                (id, tenant_id, project_id, event_type, aggregate_type, aggregate_id,
                 safe_payload, occurred_at, publish_status, publish_attempts, expires_at,
                 correlation_id, created_at, updated_at)
            VALUES (gen_random_uuid(),$1::uuid,$2::uuid,$3::varchar,
                    $6::varchar,$4::uuid,json_build_object('status',$5::varchar),
                    now(),'pending',0,now()+interval '24 hours',
                    'knowledge-worker:'||$4::uuid::text,now(),now())
            """,
            tenant_id,
            project_id,
            event_type,
            aggregate_id,
            status,
            aggregate_type,
        )
